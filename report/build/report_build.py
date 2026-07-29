"""Build the Team Binary FYP Final Report as a .docx from structured blocks."""
import copy
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import content_a
import content_b
import content_c

# Generated from the measured results by `python -m research.chapters`.
# Optional: the report still builds without it, so a checkout that has not run
# `make research` yet is not blocked — it simply omits the research chapters.
try:
    import content_research
except ImportError:                                    # not generated yet
    content_research = None

OUT = "/Users/arkamzakir/Documents/Research/Research/report/Final_Report_TeamBinary_AI_Powered_Digital_Marketing_Orchestration.docx"

BODY_FONT = "Times New Roman"
MONO_FONT = "Consolas"


# --------------------------------------------------------------- low-level
def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), v)
    return e


def add_field(paragraph, instr, placeholder="Right-click and choose “Update Field”."):
    """Insert a Word field code (TOC, PAGE, ...) into a paragraph."""
    r1 = paragraph.add_run()
    r1._r.append(_el("w:fldChar", **{"w:fldCharType": "begin"}))
    r2 = paragraph.add_run()
    t = _el("w:instrText", **{"xml:space": "preserve"})
    t.text = instr
    r2._r.append(t)
    r3 = paragraph.add_run()
    r3._r.append(_el("w:fldChar", **{"w:fldCharType": "separate"}))
    r4 = paragraph.add_run(placeholder)
    r4.italic = True
    r4.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    r5 = paragraph.add_run()
    r5._r.append(_el("w:fldChar", **{"w:fldCharType": "end"}))


def set_page_numbering(section, fmt, start=None):
    sectPr = section._sectPr
    for old in sectPr.findall(qn("w:pgNumType")):
        sectPr.remove(old)
    attrs = {"w:fmt": fmt}
    if start is not None:
        attrs["w:start"] = str(start)
    sectPr.append(_el("w:pgNumType", **attrs))


def shade(cell, hexcolor):
    cell._tc.get_or_add_tcPr().append(_el("w:shd", **{
        "w:val": "clear", "w:color": "auto", "w:fill": hexcolor}))


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(_el("w:tblHeader", **{"w:val": "true"}))


def keep_with_next(paragraph, on=True):
    paragraph.paragraph_format.keep_with_next = on


# --------------------------------------------------------------- styles
def build_styles(doc):
    st = doc.styles

    normal = st["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(6)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    specs = {
        "Heading 1": (16, True, False, RGBColor(0, 0, 0), 0, 14),
        "Heading 2": (14, True, False, RGBColor(0, 0, 0), 14, 8),
        "Heading 3": (12.5, True, False, RGBColor(0, 0, 0), 10, 6),
        "Heading 4": (12, True, True, RGBColor(0, 0, 0), 8, 4),
    }
    for name, (size, bold, italic, color, before, after) in specs.items():
        s = st[name]
        s.font.name = BODY_FONT
        s.font.size = Pt(size)
        s.font.bold = bold
        s.font.italic = italic
        s.font.color.rgb = color
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.line_spacing = 1.15
        s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        s.paragraph_format.keep_with_next = True

    cap = st["Caption"]
    cap.font.name = BODY_FONT
    cap.font.size = Pt(10)
    cap.font.italic = True
    cap.font.bold = False
    cap.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(4)
    cap.paragraph_format.space_after = Pt(12)
    cap.paragraph_format.line_spacing = 1.0

    tcap = st.add_style("Table Caption", 1)
    tcap.base_style = st["Caption"]
    tcap.font.name = BODY_FONT
    tcap.font.size = Pt(10)
    tcap.font.italic = True
    tcap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tcap.paragraph_format.space_before = Pt(10)
    tcap.paragraph_format.space_after = Pt(4)
    tcap.paragraph_format.keep_with_next = True

    code = st.add_style("Code Block", 1)
    code.font.name = MONO_FONT
    code.font.size = Pt(9)
    code.paragraph_format.line_spacing = 1.0
    code.paragraph_format.space_before = Pt(6)
    code.paragraph_format.space_after = Pt(10)
    code.paragraph_format.left_indent = Inches(0.3)
    code.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    quote = st.add_style("Note Box", 1)
    quote.font.name = BODY_FONT
    quote.font.size = Pt(11)
    quote.font.italic = True
    quote.paragraph_format.left_indent = Inches(0.35)
    quote.paragraph_format.right_indent = Inches(0.25)
    quote.paragraph_format.space_before = Pt(8)
    quote.paragraph_format.space_after = Pt(10)
    quote.paragraph_format.line_spacing = 1.3
    return doc


# --------------------------------------------------------------- helpers
def para(doc, text="", style=None, align=None, size=None, bold=False,
         italic=False, space_after=None, indent=None):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        if size:
            r.font.size = Pt(size)
    if align is not None:
        p.alignment = align
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    if indent is not None:
        p.paragraph_format.left_indent = Inches(indent)
    return p


def rich(doc, text, style=None, align=None):
    """Paragraph supporting **bold** and *italic* inline markers."""
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    import re
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            p.add_run(part[1:-1]).italic = True
        else:
            p.add_run(part)
    return p


def add_table(doc, spec):
    caption = spec.get("caption")
    if caption:
        para(doc, caption, style="Table Caption")
    header, rows = spec["header"], spec["rows"]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    hdr = t.rows[0]
    repeat_header(hdr)
    for i, h in enumerate(header):
        c = hdr.cells[i]
        c.text = ""
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(str(h))
        r.bold = True
        r.font.size = Pt(10)
        shade(c, "E8EEF6")
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            c = cells[i]
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.0
            p.alignment = (WD_ALIGN_PARAGRAPH.LEFT if i == 0
                           else WD_ALIGN_PARAGRAPH.CENTER)
            txt = str(val)
            bold = txt.startswith("**") and txt.endswith("**")
            if bold:
                txt = txt[2:-2]
            r = p.add_run(txt)
            r.font.size = Pt(10)
            r.bold = bold
    para(doc, "", space_after=6)
    return t


def add_figure(doc, spec):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(spec["path"], width=Inches(spec.get("width", 6.0)))
    para(doc, spec["caption"], style="Caption")


def add_bullets(doc, items, numbered=False):
    style = "List Number" if numbered else "List Bullet"
    for it in items:
        p = rich(doc, it, style=style)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.3
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def add_code(doc, text):
    for line in text.split("\n"):
        p = doc.add_paragraph(style="Code Block")
        p.add_run(line if line else " ")
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
    para(doc, "", space_after=8)


# --------------------------------------------------------------- render
def render(doc, blocks):
    for kind, payload in blocks:
        if kind == "h1":
            doc.add_page_break()
            doc.add_heading(payload, level=1)
        elif kind == "h1_nobreak":
            doc.add_heading(payload, level=1)
        elif kind == "h2":
            doc.add_heading(payload, level=2)
        elif kind == "h3":
            doc.add_heading(payload, level=3)
        elif kind == "h4":
            doc.add_heading(payload, level=4)
        elif kind == "p":
            rich(doc, payload)
        elif kind == "note":
            rich(doc, payload, style="Note Box")
        elif kind == "bullets":
            add_bullets(doc, payload)
        elif kind == "numbers":
            add_bullets(doc, payload, numbered=True)
        elif kind == "table":
            add_table(doc, payload)
        elif kind == "figure":
            add_figure(doc, payload)
        elif kind == "code":
            add_code(doc, payload)
        elif kind == "pagebreak":
            doc.add_page_break()
        else:
            raise ValueError(f"unknown block: {kind}")


# --------------------------------------------------------------- front pages
def title_pages(doc):
    def centred(text, size, bold=False, after=10, before=0, italic=False):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.line_spacing = 1.2
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        r.font.size = Pt(size)
        return p

    centred("FINAL REPORT", 22, True, after=18, before=36)
    centred("Level 04", 16, True, after=26)
    centred("AI-Powered Digital Marketing Orchestration", 18, True, after=8)
    centred("An Integrated, Closed-Loop Framework for Cold-Start Digital Marketing",
            12, False, after=40, italic=True)
    centred("Team Binary", 15, True, after=26)

    t = doc.add_table(rows=4, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (idx, name) in enumerate(MEMBERS):
        for j, val in enumerate((idx, name)):
            c = t.cell(i, j)
            c.text = ""
            p = c.paragraphs[0]
            p.alignment = (WD_ALIGN_PARAGRAPH.RIGHT if j == 0
                           else WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.2
            r = p.add_run(val)
            r.font.size = Pt(12)
    para(doc, "", space_after=30)

    centred("Supervised by", 12, False, after=6)
    centred("Dr. A. L. A. Romesh R. Thanuja   (IN/CM Supervisor)", 12, True, after=4)
    centred("Ms. M. A. N. Perera   (IDS Supervisor)", 12, True, after=34)
    centred("Faculty of Information Technology", 13, False, after=4)
    centred("University of Moratuwa", 13, False, after=4)
    centred("July 2026", 13, False, after=0)

    doc.add_page_break()

    centred("AI-Powered Digital Marketing Orchestration", 18, True,
            after=10, before=60)
    centred("An Integrated, Closed-Loop Framework for Cold-Start Digital Marketing",
            12, False, after=48, italic=True)
    centred("Team Binary", 15, True, after=24)
    for idx, name in MEMBERS:
        centred(f"{idx}     {name}", 12, False, after=4)
    para(doc, "", space_after=40)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(46)
    r = p.add_run("Dissertation submitted to the Faculty of Information Technology, "
                  "University of Moratuwa, Sri Lanka, for the partial fulfilment of "
                  "the requirements of the Degree of Bachelor of Science Honours in "
                  "Information Technology.")
    r.font.size = Pt(12)
    centred("July 2026", 13, False, after=0)


MEMBERS = [
    ("215001G", "Aadhil M. H. M."),
    ("215015D", "Arqam Z. H."),
    ("215110N", "Sarah M. M. F."),
    ("215129F", "Zanar M. H. M. R. A."),
]


def declaration(doc):
    doc.add_page_break()
    doc.add_heading("DECLARATION", level=1)
    rich(doc, "We declare that this thesis is our own work and has not been "
              "submitted in any form for another degree or diploma at any "
              "university or other institution of tertiary education. Information "
              "derived from the published or unpublished work of others has been "
              "acknowledged in the text and a list of references is given.")
    para(doc, "", space_after=22)
    for idx, name in MEMBERS:
        p = para(doc, f"………………………………………          {name}   ({idx})")
        p.paragraph_format.space_after = Pt(20)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para(doc, "Date: ……… / ……… / 2026", space_after=30)
    para(doc, "Supervised by", bold=True, space_after=18)
    para(doc, "………………………………………          Dr. A. L. A. Romesh R. Thanuja",
         space_after=6)
    para(doc, "Date: ……… / ……… / 2026", space_after=24)
    para(doc, "………………………………………          Ms. M. A. N. Perera", space_after=6)
    para(doc, "Date: ……… / ……… / 2026", space_after=6)


def toc_pages(doc):
    doc.add_page_break()
    doc.add_heading("TABLE OF CONTENTS", level=1)
    p = doc.add_paragraph()
    add_field(p, ' TOC \\o "1-3" \\h \\z \\u ',
              "Table of contents — place the cursor here, press F9 and choose "
              "“Update entire table”.")

    doc.add_page_break()
    doc.add_heading("LIST OF FIGURES", level=1)
    p = doc.add_paragraph()
    add_field(p, ' TOC \\h \\z \\t "Caption,1" ',
              "List of figures — place the cursor here and press F9.")

    doc.add_page_break()
    doc.add_heading("LIST OF TABLES", level=1)
    p = doc.add_paragraph()
    add_field(p, ' TOC \\h \\z \\t "Table Caption,1" ',
              "List of tables — place the cursor here and press F9.")


def add_footer_pagenum(section, italic_label=None):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(p, " PAGE ", "1")
    for r in p.runs:
        r.font.size = Pt(10)
        r.font.name = BODY_FONT
        r.italic = False
        r.font.color.rgb = RGBColor(0, 0, 0)


# --------------------------------------------------------------- main
def main():
    doc = Document()
    build_styles(doc)

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)   # A4
    sec.left_margin = Inches(1.25)
    sec.right_margin = Inches(1.0)
    sec.top_margin = Inches(1.0)
    sec.bottom_margin = Inches(1.0)
    sec.different_first_page_header_footer = True
    set_page_numbering(sec, "lowerRoman", start=1)

    title_pages(doc)
    render(doc, content_a.FRONT_BLOCKS)
    declaration(doc)
    toc_pages(doc)
    add_footer_pagenum(sec)

    body = doc.add_section(WD_SECTION.NEW_PAGE)
    body.page_width, body.page_height = Inches(8.27), Inches(11.69)
    body.left_margin = Inches(1.25)
    body.right_margin = Inches(1.0)
    body.top_margin = Inches(1.0)
    body.bottom_margin = Inches(1.0)
    body.different_first_page_header_footer = False
    body.footer.is_linked_to_previous = False
    set_page_numbering(body, "decimal", start=1)
    add_footer_pagenum(body)

    render(doc, content_a.BODY_BLOCKS)
    render(doc, content_b.BLOCKS)
    render(doc, content_c.BLOCKS)
    if content_research is not None:
        render(doc, content_research.BLOCKS)
    else:
        print("note: research chapters omitted — run `make research` then "
              "`python -m research.chapters` to include them.")

    doc.save(OUT)
    print("saved:", OUT)


if __name__ == "__main__":
    main()
